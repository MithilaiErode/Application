"""Offline tests: parser, output writers and the web flow with the Claude call stubbed out."""

import io
import time
import zipfile
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import analyzer, main
from app.parser import parse_upload

SAMPLES = Path(__file__).parent.parent / "samples" / "classic"


def _read(name: str) -> bytes:
    return (SAMPLES / name).read_bytes()


def test_parses_escaped_xml_content():
    [rec] = parse_upload("SalesOrderEntry_SO_Custom_v3.xml", _read("SalesOrderEntry_SO_Custom_v3.xml"))
    assert rec.name == "SO_Custom_v3"
    assert rec.form == "SalesOrderEntry.SalesOrderForm"
    assert "BeforeAdapterMethod" in rec.script
    assert "epiTextBoxReference" in rec.ui_context
    assert "[C# SCRIPT EXTRACTED SEPARATELY]" in rec.ui_context


def test_parses_base64_content():
    [rec] = parse_upload("JobEntry_JobShipDate.xml", _read("JobEntry_JobShipDate.xml"))
    assert rec.name == "JobShipDate"
    assert "ReqDueDate" in rec.script
    assert "epiLabelLate" in rec.ui_context


def test_parses_cs_and_zip():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for f in SAMPLES.iterdir():
            zf.writestr(f"exports/{f.name}", f.read_bytes())
    recs = parse_upload("all.zip", buf.getvalue())
    assert sorted(r.name for r in recs) == ["JobShipDate", "PartMaint_PartExtra_old", "SO_Custom_v3"]


def test_multiple_rows_in_one_file():
    xml = b"""<Rows>
      <XXXDef><TypeCode>Customization</TypeCode><Key1>A</Key1><Key2>F1</Key2>
        <Content>public class Script { void InitializeCustomCode() { var x = oTrans; } }</Content></XXXDef>
      <XXXDef><TypeCode>Personalization</TypeCode><Key1>P</Key1><Key2>F1</Key2><Content>&lt;x/&gt;</Content></XXXDef>
      <XXXDef><TypeCode>Customization</TypeCode><Key1>B</Key1><Key2>F2</Key2>
        <Content>public class Script { void InitializeCustomCode() { var y = oTrans; } }</Content></XXXDef>
    </Rows>"""
    recs = parse_upload("extract.xml", xml)
    assert [r.name for r in recs] == ["A", "B"]


FAKE = {
    "what_it_does": ["Hides Reference", "Defaults Ship Via", "Requires PO"],
    "parts": [{"behaviour": "Requires PO", "classic_mechanism": "BeforeAdapterMethod",
               "kinetic_approach": "BPM", "reason": "Server-side validation"}],
    "bucket": "Layer + BPM/Function", "confidence": 88, "confidence_reason": "Clear",
    "hours_build": 3, "hours_test": 2, "hours_deploy": 1, "estimate_notes": "n/a",
    "improvements": ["Closes DMT gap"], "risks": [], "questions_for_client": ["Is EXP still used?"],
    "files": [{"path": "bpm/../../evil.cs", "file_type": "bpm_code", "description": "BPM", "content": "// code"}],
    "import_steps": ["Create directive"],
    "test_cases": [{"test": "Save without PO", "expected": "Error", "channel": "Screen"}],
}


def test_web_flow(monkeypatch, tmp_path):
    monkeypatch.setattr(main, "JOBS_DIR", tmp_path)
    monkeypatch.setattr(main, "APP_PASSWORD", "")
    monkeypatch.setattr(analyzer, "analyse", lambda item, prompt: (FAKE, {
        "model": "stub", "input_tokens": 10, "output_tokens": 5, "cache_read_tokens": 0, "cache_write_tokens": 0}))
    client = TestClient(main.app)
    files = [("files", (f.name, f.read_bytes())) for f in sorted(SAMPLES.iterdir())]
    res = client.post("/api/jobs", data={"client": "ClientX", "hourly_rate": "180"}, files=files)
    assert res.status_code == 200, res.text
    job_id = res.json()["id"]
    for _ in range(50):
        job = client.get(f"/api/jobs/{job_id}").json()
        if job["status"] != "running":
            break
        time.sleep(0.1)
    assert job["status"] == "completed"
    assert job["done"] == 3

    item_id = job["items"][0]["id"]
    detail = client.get(f"/api/jobs/{job_id}/items/{item_id}").json()
    assert "Layer + BPM/Function" in detail["analysis_md"]

    z = client.get(f"/api/jobs/{job_id}/download")
    assert z.status_code == 200
    names = zipfile.ZipFile(io.BytesIO(z.content)).namelist()
    assert "inventory.xlsx" in names and "client_report.docx" in names
    assert not any(".." in n for n in names)
    assert any(n.endswith("bpm/evil.cs") for n in names)  # path traversal neutralised
    assert any(n.endswith("TEST_STEPS.md") for n in names)

    wb = load_workbook(io.BytesIO(zipfile.ZipFile(io.BytesIO(z.content)).read("inventory.xlsx")))
    assert wb.sheetnames == ["Summary", "Inventory"]
    assert wb["Inventory"]["J2"].value == "=SUM(G2:I2)"


def _run(client, files, **form):
    res = client.post("/api/jobs", data={"client": "ClientX", **form}, files=files)
    assert res.status_code == 200, res.text
    job_id = res.json()["id"]
    for _ in range(100):
        job = client.get(f"/api/jobs/{job_id}").json()
        if job["status"] != "running":
            return job
        time.sleep(0.05)
    raise AssertionError("job did not finish")


def _counting_stub(monkeypatch, tmp_path):
    calls = []
    monkeypatch.setattr(main, "JOBS_DIR", tmp_path)
    monkeypatch.setattr(main, "APP_PASSWORD", "")

    def fake(item, prompt):
        calls.append(item.name)
        return FAKE, {"model": "claude-sonnet-5", "input_tokens": 1000, "output_tokens": 10000,
                      "cache_read_tokens": 0, "cache_write_tokens": 0}
    monkeypatch.setattr(analyzer, "analyse", fake)
    return calls


def test_same_code_in_one_job_is_analysed_once(monkeypatch, tmp_path):
    calls = _counting_stub(monkeypatch, tmp_path)
    data = _read("SalesOrderEntry_SO_Custom_v3.xml")
    job = _run(TestClient(main.app), [("files", ("companyA.xml", data)), ("files", ("companyB.xml", data))])
    assert job["status"] == "completed"
    assert len(calls) == 1
    assert [i["status"] for i in job["items"]] == ["done", "done"]
    assert job["items"][1]["reused"].startswith("same code as")
    assert job["cost_usd"] == 0.102  # one Sonnet 5 call: 1k input @ $2/M + 10k output @ $10/M


def test_rerun_reuses_stored_results(monkeypatch, tmp_path):
    calls = _counting_stub(monkeypatch, tmp_path)
    client = TestClient(main.app)
    files = [("files", (f.name, f.read_bytes())) for f in sorted(SAMPLES.iterdir())]
    first = _run(client, files, hourly_rate="180")
    assert len(calls) == 3 and first["cost_usd"] > 0
    second = _run(client, files, hourly_rate="180")
    assert len(calls) == 3, "identical re-run must not call the AI again"
    assert second["status"] == "completed" and second["cost_usd"] == 0
    assert all(i["reused"] == "stored result" for i in second["items"])
    third = _run(client, files, hourly_rate="200")  # different client profile → fresh analysis
    assert len(calls) == 6 and third["cost_usd"] > 0


def test_clashing_generated_paths_are_kept(tmp_path):
    from app import outputs
    item = {"id": "x", "name": "N", "form": "F", "source_file": "s", "script": "", "warnings": [],
            "analysis": {**FAKE, "files": [
                {"path": "bpm", "file_type": "other", "description": "", "content": "A"},
                {"path": "bpm/x.cs", "file_type": "bpm_code", "description": "", "content": "B"}]}}
    folder = outputs.write_item_folder(tmp_path, item)
    assert (folder / "bpm").read_text() == "A"
    assert (folder / "other_files" / "02_bpm_x.cs").read_text() == "B"


def test_same_code_on_different_forms_is_not_merged(monkeypatch, tmp_path):
    calls = _counting_stub(monkeypatch, tmp_path)
    data = _read("SalesOrderEntry_SO_Custom_v3.xml")
    other_form = data.replace(b"SalesOrderEntry.SalesOrderForm", b"QuoteEntry.QuoteForm")
    job = _run(TestClient(main.app), [("files", ("a.xml", data)), ("files", ("b.xml", other_form))])
    assert len(calls) == 2 and not any(i["reused"] for i in job["items"])


def test_login_required(monkeypatch):
    monkeypatch.setattr(main, "APP_USER", "mithilai")
    monkeypatch.setattr(main, "APP_PASSWORD", "secret")
    client = TestClient(main.app)
    assert client.get("/").status_code == 401
    assert client.get("/", auth=("mithilai", "secret")).status_code == 200
    assert client.get("/api/health").status_code == 200
