"""Write per-customisation folders, the Excel inventory, the client report and the zip."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path, PurePosixPath

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .analyzer import BUCKETS

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def safe_name(text: str, fallback: str = "item") -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", text).strip("._")
    return cleaned[:80] or fallback


def _safe_rel_path(path: str) -> PurePosixPath:
    parts = [safe_name(p) for p in PurePosixPath(path.replace("\\", "/")).parts if p not in ("", ".", "..", "/")]
    return PurePosixPath(*parts) if parts else PurePosixPath("file.txt")


def total_hours(a: dict) -> float:
    return round(float(a.get("hours_build", 0)) + float(a.get("hours_test", 0)) + float(a.get("hours_deploy", 0)), 1)


def _bullets(items: list[str]) -> str:
    return "\n".join(f"- {i}" for i in items) if items else "- None"


def analysis_markdown(item: dict) -> str:
    a = item["analysis"]
    parts = "\n".join(
        f"| {p['behaviour']} | {p['classic_mechanism']} | {p['kinetic_approach']} | {p['reason']} |"
        for p in a["parts"]
    )
    files = "\n".join(f"- `{f['path']}` – {f['description']}" for f in a["files"]) or "- None"
    warnings = _bullets(item.get("warnings", []))
    return f"""# {item['name']}

**Form:** {item['form'] or 'unknown'}
**Source file:** {item['source_file']}
**Status:** DRAFT – consultant review required

## What it does
{_bullets(a['what_it_does'])}

## Decision
| Behaviour | Classic mechanism | Kinetic approach | Why |
|---|---|---|---|
{parts}

**Bucket:** {a['bucket']} · **Confidence:** {a['confidence']}% – {a['confidence_reason']}

## Estimate
Build {a['hours_build']} h · Test {a['hours_test']} h · Deploy {a['hours_deploy']} h · **Total {total_hours(a)} h**

{a['estimate_notes']}

## Improvements
{_bullets(a['improvements'])}

## Risks
{_bullets(a['risks'])}

## Questions for the client
{_bullets(a['questions_for_client'])}

## Generated files
{files}

## Parser warnings
{warnings}
"""


def test_steps_markdown(item: dict) -> str:
    a = item["analysis"]
    steps = "\n".join(f"{i}. {s}" for i, s in enumerate(a["import_steps"], 1)) or "1. Nothing to import."
    rows = "\n".join(
        f"| {i} | {t['test']} | {t['channel']} | {t['expected']} |" for i, t in enumerate(a["test_cases"], 1)
    )
    return f"""# Import & test – {item['name']}

> Import into the **TEST / PILOT** environment only. Move to production through a
> Solution Workbench package after testing and consultant sign-off.

## Import steps
{steps}

## Test cases
| # | Test | Channel | Expected |
|---|---|---|---|
{rows}
"""


def write_item_folder(base: Path, item: dict) -> Path:
    folder = base / safe_name(f"{item['form'] or 'Form'}_{item['name']}", item["id"])
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "ANALYSIS.md").write_text(analysis_markdown(item), encoding="utf-8")
    (folder / "TEST_STEPS.md").write_text(test_steps_markdown(item), encoding="utf-8")
    if item.get("script"):
        (folder / "original_script.cs").write_text(item["script"], encoding="utf-8")
    for n, f in enumerate(item["analysis"]["files"], 1):
        target = folder / Path(*_safe_rel_path(f["path"]).parts)
        try:
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(f["content"], encoding="utf-8")
        except OSError:
            # Clashing generated paths (e.g. "bpm" and "bpm/x.cs"): keep the file under a flat name.
            fallback = folder / "other_files" / f"{n:02d}_{safe_name(f['path'], 'file.txt')}"
            fallback.parent.mkdir(parents=True, exist_ok=True)
            fallback.write_text(f["content"], encoding="utf-8")
    return folder


def _style_header(ws, row: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        c = ws.cell(row=row, column=col)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER


def write_inventory(path: Path, job: dict) -> None:
    items = [i for i in job["items"] if i.get("analysis")]
    wb = Workbook()
    inv = wb.active
    inv.title = "Inventory"
    headers = ["#", "Form", "Customisation", "What it does", "Bucket", "Confidence %",
               "Build h", "Test h", "Deploy h", "Total h", "Files generated", "Risks", "Status"]
    inv.append(headers)
    _style_header(inv, 1, len(headers))
    for n, item in enumerate(items, 1):
        a = item["analysis"]
        r = n + 1
        inv.append([
            n, item["form"], item["name"], "; ".join(a["what_it_does"]), a["bucket"], a["confidence"],
            a["hours_build"], a["hours_test"], a["hours_deploy"], f"=SUM(G{r}:I{r})",
            ", ".join(f["path"] for f in a["files"]) or "-", "; ".join(a["risks"]) or "-",
            "Draft – review required",
        ])
    failed = [i for i in job["items"] if i.get("error")]
    for n, item in enumerate(failed, len(items) + 1):
        inv.append([n, item["form"], item["name"], "", "Needs Review", "", "", "", "", 0,
                    "", item["error"], "Analysis failed"])
    last = max(2, inv.max_row)
    for row in inv.iter_rows(min_row=2, max_row=inv.max_row):
        for c in row:
            c.font = Font(name=FONT)
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for col, width in zip(range(1, 14), [5, 22, 26, 60, 18, 12, 9, 9, 9, 9, 40, 50, 22]):
        inv.column_dimensions[get_column_letter(col)].width = width
    inv.freeze_panes = "D2"
    inv.auto_filter.ref = f"A1:M{last}"

    s = wb.create_sheet("Summary")
    s["A1"] = f"Kinetic Migration Assessment – {job['profile'].get('Client') or 'Client'}"
    s["A1"].font = Font(name=FONT, bold=True, size=14, color="1F3864")
    s["A3"], s["B3"] = "Hourly rate", float(job["profile"].get("Hourly rate") or 0)
    s["C3"] = "Input – from the client profile entered in the app; edit to recalculate."
    s["B3"].fill = INPUT_FILL
    s["B3"].font = Font(name=FONT, color="0000FF")
    s["B3"].number_format = "$#,##0"
    s["A5"], s["B5"], s["C5"], s["D5"] = "Bucket", "Customisations", "Hours", "Cost"
    _style_header(s, 5, 4)
    rng_b, rng_h = f"Inventory!$E$2:$E${last}", f"Inventory!$J$2:$J${last}"
    for i, bucket in enumerate(BUCKETS, start=6):
        s.cell(row=i, column=1, value=bucket)
        s.cell(row=i, column=2, value=f'=COUNTIF({rng_b},A{i})')
        s.cell(row=i, column=3, value=f'=SUMIF({rng_b},A{i},{rng_h})')
        s.cell(row=i, column=4, value=f"=C{i}*$B$3").number_format = "$#,##0"
    t = 6 + len(BUCKETS)
    s.cell(row=t, column=1, value="Total")
    s.cell(row=t, column=2, value=f"=SUM(B6:B{t - 1})")
    s.cell(row=t, column=3, value=f"=SUM(C6:C{t - 1})")
    s.cell(row=t, column=4, value=f"=SUM(D6:D{t - 1})").number_format = "$#,##0"
    for row in s.iter_rows(min_row=6, max_row=t):
        for c in row:
            c.font = Font(name=FONT, bold=(c.row == t))
            c.border = BORDER
    s["A3"].font = s["C3"].font = Font(name=FONT)
    s.cell(row=t + 2, column=1, value="Estimates are AI drafts for consultant review, not a quote.").font = \
        Font(name=FONT, italic=True, color="7F7F7F")
    for col, width in zip("ABCD", [24, 16, 12, 14]):
        s.column_dimensions[col].width = width
    wb.move_sheet("Summary", offset=-1)
    wb.active = 0
    wb.save(path)


def write_client_report(path: Path, job: dict) -> None:
    items = [i for i in job["items"] if i.get("analysis")]
    profile = job["profile"]
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(10)
    doc.add_heading(f"Kinetic Migration Assessment – {profile.get('Client') or 'Client'}", 0)
    doc.add_paragraph(
        f"Current system: {profile.get('Current version') or 'n/a'}    "
        f"Target: {profile.get('Target version') or 'n/a'}"
    )
    doc.add_paragraph("DRAFT – prepared with the Mithilai Migration Assessor and subject to consultant review.")

    counts = {b: sum(1 for i in items if i["analysis"]["bucket"] == b) for b in BUCKETS}
    hours = sum(total_hours(i["analysis"]) for i in items)
    doc.add_heading("Executive summary", 1)
    doc.add_paragraph(
        f"{len(job['items'])} customisations assessed ({len(items)} analysed). "
        + ", ".join(f"{n} {b}" for b, n in counts.items() if n)
        + f". Estimated effort: {hours:,.0f} hours."
    )
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Bucket", "Customisations"
    for b, n in counts.items():
        row = table.add_row().cells
        row[0].text, row[1].text = b, str(n)

    doc.add_heading("Customisations", 1)
    t = doc.add_table(rows=1, cols=5)
    t.style = "Light Grid Accent 1"
    for c, h in zip(t.rows[0].cells, ["Form", "Customisation", "Approach", "Confidence", "Hours"]):
        c.text = h
    for i in items:
        a = i["analysis"]
        row = t.add_row().cells
        row[0].text, row[1].text, row[2].text = i["form"] or "-", i["name"], a["bucket"]
        row[3].text, row[4].text = f"{a['confidence']}%", f"{total_hours(a)}"

    for title, key in [("Improvements found", "improvements"), ("Risks", "risks"),
                       ("Decisions and questions for the client", "questions_for_client")]:
        doc.add_heading(title, 1)
        entries = [(i["name"], e) for i in items for e in i["analysis"][key]]
        if not entries:
            doc.add_paragraph("None identified.")
        for name, e in entries:
            doc.add_paragraph(f"{name}: {e}", style="List Bullet")

    doc.add_heading("Next steps", 1)
    for step in ["Consultant review of all drafts and estimates.",
                 "Answer the open questions above.",
                 "Import approved changes into the TEST environment and run the provided test cases.",
                 "Package approved changes with Solution Workbench and deploy to production."]:
        doc.add_paragraph(step, style="List Number")
    doc.save(path)


def build_zip(job_dir: Path, zip_path: Path) -> None:
    out = job_dir / "output"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for file in sorted(out.rglob("*")):
            if file.is_file():
                zf.write(file, file.relative_to(out))
